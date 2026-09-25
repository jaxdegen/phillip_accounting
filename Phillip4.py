import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os
import subprocess
import ollama
import json
from datetime import datetime

expense_words = [
    "bought",
    "buy",
    "buying",
    "purchased",
    "purchase",
    "paid",
    "pay",
    "paying",
    "spent",
    "spend",
    "spending",
    "invested",
    "investing",
    "investment",
    "got",
    "obtained",
    "obtaining",
    "obtain",
    "financed",
    "financing",
    "to finance",
    "repurchased",
    "repurchasing"
]

income_words = [
    "sold",
    "sell",
    "selling",
    "received",
    "receive",
    "receiving",
    "earned",
    "earn",
    "earning",
    "collected",
    "collect",
    "collecting",
    "distribute",
    "distributed",
    "distributing"
]

# The spreadsheet Phillip is currently using
current_file = "accounting.xlsx"

ai_name = "Phillip"


def create_new_spreadsheet():
    global current_file

    filename = new_file_entry.get().strip()

    if not filename:
        messagebox.showwarning(
            "Phillip",
            "Please enter a name for the spreadsheet."
        )
        return

    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Account Log"

    # Create the accounting columns
    sheet.append([
        "Date",
        "Vendor/Client",
        "Description",
        "Quantity",
        "Category",
        "Type",
        "Amount"
    ])

    workbook.save(filename)

    # Make the new spreadsheet Phillip's current spreadsheet
    current_file = os.path.abspath(filename)

    messagebox.showinfo(
        "Phillip",
        f"New spreadsheet created:\n\n{filename}"
    )

    new_file_entry.delete(0, tk.END)

    status_label.config(
        text=f"Current spreadsheet: {filename}",
        font=("Times New Roman", 11)
    )

    # Refresh spreadsheet list
    refresh_spreadsheet_list()


def add_transaction():
    global current_file

    transaction = entry.get().strip()

    if not transaction:
        messagebox.showwarning(
            "Phillip",
            "Please enter a transaction."
        )
        return

    try:
        response = ollama.chat(
            model="llama3.2:3b",
            messages=[
                {
                    "role": "system",
                    "content": """
You are Phillip, an accounting assistant.

Analyze the user's transaction and return ONLY valid JSON.

The JSON must contain exactly these fields:

vendor
description
quantity
category
type
amount

Rules:
- vendor = company, person, or customer involved
- description = what was purchased, sold, or paid for
- quantity = number of items involved, or 1 if a quantity is not stated
- category = ALWAYS provide an appropriate accounting category. Never leave category blank.- type = ONLY "Income" or "Expense"
- type = ONLY "Income" or "Expense"
- amount = numerical amount only
- Do not include dollar signs in amount
- Do not include explanations outside the JSON
"""
                },
                {
                    "role": "user",
                    "content": transaction
                }
            ],
            format="json"
        )

        ai_text = response["message"]["content"]

        # Convert the JSON response into Python data
        data = json.loads(ai_text)

        vendor = data["vendor"]
        if not vendor:
            vendor = "N/A"
        description = data["description"]
        if not description:
             description = "N/A"
        quantity = data["quantity"]
        if not description:
            quantity = "1"
        category = data["category"]
        if not category:
            category = "Uncategorized"
        transaction_lower = transaction.lower()

        if any(word in transaction_lower for word in expense_words):
            transaction_type = "Expense"

        elif any(word in transaction_lower for word in income_words):
            transaction_type = "Income"

        else:
            transaction_type = data["type"]

        if transaction_type not in ["Income", "Expense"]:
            transaction_type = "Expense"
            
        amount = float(data["amount"])
        

    except Exception as e:
        messagebox.showerror(
            "Phillip Error",
            f"Phillip could not process the transaction:\n\n{e}"
        )
        return

    # Create the spreadsheet if it doesn't exist
    if os.path.exists(current_file):
        workbook = load_workbook(current_file)
        sheet = workbook["Account Log"]

    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Account Log"

        # Create headers
        sheet.append([
            "Date",
            "Vendor/Client",
            "Description",
            "Quantity",
            "Category",
            "Type",
            "Amount"
        ])

    # Automatically get today's date
    date = datetime.now().strftime("%m-%d-%Y")

    # Add Phillip's categorized transaction
    sheet.append([
        date,
        vendor,
        description,
        quantity,
        category,
        transaction_type,
        amount
    ])

    # Format the Amount cell as currency
    amount_cell = sheet.cell(row=sheet.max_row, column=7)
    amount_cell.number_format = '$#,##0.00'

    workbook.save(current_file)

    # Open the spreadsheet
    subprocess.run(["open", current_file])

    status_label.config(
        text=f"Data successfully added to {os.path.basename(current_file)}!",
        font=("Times New Roman", 11)
    )

    entry.delete(0, tk.END)


# -------------------------
# SEARCH SPREADSHEETS
# -------------------------

def refresh_spreadsheet_list():
    spreadsheet_list.delete(0, tk.END)

    search_text = search_entry.get().lower()

    folder = os.path.dirname(os.path.abspath(__file__))

    for filename in os.listdir(folder):
        if filename.endswith(".xlsx"):
            if search_text in filename.lower():
                spreadsheet_list.insert(tk.END, filename)


def search_spreadsheets(event=None):
    refresh_spreadsheet_list()


def select_spreadsheet():
    global current_file

    selection = spreadsheet_list.curselection()

    if not selection:
        messagebox.showwarning(
            "Phillip",
            "Please select a spreadsheet."
        )
        return

    filename = spreadsheet_list.get(selection[0])

    current_file = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        filename
    )

    status_label.config(
        text=f"Current spreadsheet: {filename}",
        font=("Times New Roman", 11)
    )

def delete_spreadsheet():
    global current_file

    selection = spreadsheet_list.curselection()

    if not selection:
        messagebox.showwarning(
            "Phillip",
            "Please select a spreadsheet to delete."
        )
        return

    filename = spreadsheet_list.get(selection[0])

    confirm = messagebox.askyesno(
        "Delete Spreadsheet",
        f"Are you sure you want to permanently delete:\n\n{filename}?"
    )

    if not confirm:
        return

    filepath = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        filename
    )

    try:
        os.remove(filepath)

        # If the deleted spreadsheet was the current spreadsheet,
        # switch back to the default spreadsheet name
        if os.path.abspath(current_file) == os.path.abspath(filepath):
            current_file = "accounting.xlsx"

        refresh_spreadsheet_list()

        status_label.config(
            text=f"Deleted spreadsheet: {filename}",
            font=("Times New Roman", 11)
        )

    except Exception as e:
        messagebox.showerror(
            "Phillip Error",
            f"Phillip could not delete the spreadsheet:\n\n{e}"
        )

# -------------------------
# WINDOW
# -------------------------

window = tk.Tk()

window.title("Phillip Accounting")
window.geometry("700x650")


title = tk.Label(
    window,
    text="PHILLIP ACCOUNTING",
    font=("Times New Roman", 23, "bold")
)

title.pack(pady=20)


instructions = tk.Label(
    window,
    text="Input New Data Here:",
    font=("Times New Roman", 15)
)

instructions.pack()


entry = tk.Entry(
    window,
    width=45,
    font=("Courier New", 12)
)

entry.pack(pady=15)


submit_button = tk.Button(
    window,
    text="Enter Data",
    font=("Times New Roman", 13, "bold"),
    command=add_transaction
)

submit_button.pack(pady=10)


# -------------------------
# SPREADSHEET SEARCH
# -------------------------

search_label = tk.Label(
    window,
    text="Search Existing Spreadsheets:",
    font=("Times New Roman", 11)
)

search_label.pack(pady=(15, 5))


search_entry = tk.Entry(
    window,
    width=40,
    font=("Courier New", 12)
)

search_entry.pack(pady=5)

search_entry.bind(
    "<KeyRelease>",
    search_spreadsheets
)


spreadsheet_list = tk.Listbox(
    window,
    width=35,
    height=3,
    font=("Courier New", 11)
)

spreadsheet_list.pack(pady=5)


select_button = tk.Button(
    window,
    text="Select Spreadsheet",
    font=("Times New Roman", 11),
    command=select_spreadsheet
)

select_button.pack(pady=5)

delete_button = tk.Button(
    window,
    text="Delete Spreadsheet",
    font=("Times New Roman", 11),
    command=delete_spreadsheet
)

delete_button.pack(pady=5)

# -------------------------
# NEW SPREADSHEET
# -------------------------

status_label = tk.Label(
    window,
    text=f"Current spreadsheet: {current_file}",
    font=("Times New Roman", 11)
)

status_label.pack(pady=10)


new_file_entry = tk.Entry(
    window,
    width=40,
    font=("Courier New", 12)
)

new_file_entry.pack(pady=5)


new_file_button = tk.Button(
    window,
    text="Create New Spreadsheet",
    font=("Courier New", 12),
    command=create_new_spreadsheet
)

new_file_button.pack(pady=5)


# -------------------------
# FOOTER
# -------------------------

footer = tk.Label(
    window,
    text="Powered by DeBeaul Tech and Solutions",
    font=("Times New Roman", 8)
)

footer.pack(pady=10)


# Load spreadsheets when Phillip starts
refresh_spreadsheet_list()


window.mainloop()
