from flask import Flask, jsonify, request
import phillip_core
import os
from pathlib import Path

app = Flask(
    __name__,
    static_folder="website",
    static_url_path=""
)


@app.route("/")
def home():
    return app.send_static_file("index.html")


@app.route("/api/templates")
def templates():
    return jsonify({
        "templates": phillip_core.template_options
    })

@app.route("/api/spreadsheets")
def spreadsheets():from flask import Flask, jsonify, request
import phillip_core
import os
from pathlib import Path

app = Flask(
    __name__,
    static_folder="website",
    static_url_path=""
)


@app.route("/")
def home():
    return app.send_static_file("index.html")


@app.route("/api/templates")
def templates():
    return jsonify({
        "templates": phillip_core.template_options
    })

@app.route("/api/spreadsheets")
def spreadsheets():from flask import Flask, jsonify, request
import phillip_core
import os
from pathlib import Path

app = Flask(
    __name__,
    static_folder="website",
    static_url_path=""
)


@app.route("/")
def home():
    return app.send_static_file("index.html")


@app.route("/api/templates")
def templates():
    return jsonify({
        "templates": phillip_core.template_options
    })

@app.route("/api/spreadsheets")
def spreadsheets():from flask import Flask, jsonify, request
import phillip_core
import os
from pathlib import Path

app = Flask(
    __name__,
    static_folder="website",
    static_url_path=""
)


@app.route("/")
def home():
    return app.send_static_file("index.html")


@app.route("/api/templates")
def templates():
    return jsonify({
        "templates": phillip_core.template_options
    })

@app.route("/api/spreadsheets")
def spreadsheets():from flask import Flask, jsonify, request
import phillip_core
import os
from pathlib import Path

app = Flask(
    __name__,
    static_folder="website",
    static_url_path=""
)


@app.route("/")
def home():
    return app.send_static_file("index.html")


@app.route("/api/templates")
def templates():
    return jsonify({
        "templates": phillip_core.template_options
    })

@app.route("/api/spreadsheets")
def spreadsheets():
    folder = Path(phillip_core.DATA_FOLDER)
    search = request.args.get("search", "").strip().lower()

    files = sorted(
        file.name
        for file in folder.glob("*.xlsx")
        if search in file.name.lower()
    )

    return jsonify({
        "spreadsheets": files
    })

@app.route("/api/transaction", methods=["POST"])
def transaction():
    data = request.get_json() or {}

    text = data.get("text", "").strip()
    template = data.get("template", "")

    if not text:
        return jsonify({
            "error": "No transaction information provided."
        }), 400

    if template not in phillip_core.template_options:
        return jsonify({
            "error": "Invalid spreadsheet template."
        }), 400

    result = phillip_core.add_transaction(
        transaction=text,
        template=template,
        web=True
    )

    return jsonify(result)


@app.route("/api/kinematics/1d", methods=["POST"])
def kinematics_1d():
    data = request.get_json() or {}

    try:
        result = phillip_core.calculate_1d_kinematics(data)

        return jsonify({
            "status": "success",
            "result": result
        })

    except Exception as error:
        return jsonify({
            "error": str(error)
        }), 400

@app.route("/api/kinematics/2d", methods=["POST"])
def kinematics_2d():
    data = request.get_json() or {}

    try:
        planet = data.get("planet", "Earth")

        result = phillip_core.calculate_2d_kinematics(
            data,
            planet=planet
        )

        return jsonify({
            "status": "success",
            "result": result
        })

    except Exception as error:
        return jsonify({
            "error": str(error)
        }), 400

@app.route("/api/select-spreadsheet", methods=["POST"])
def select_spreadsheet():
    data = request.get_json() or {}
    filename = data.get("filename", "").strip()

    filepath = os.path.join(
        phillip_core.DATA_FOLDER,
        filename
    )

    if not os.path.exists(filepath):
        return jsonify({
            "error": "Spreadsheet not found."
        }), 404

    phillip_core.current_file = filepath

    return jsonify({
        "status": "success",
        "filename": filename
    })

@app.route("/api/delete-spreadsheet", methods=["POST"])
def delete_spreadsheet():
    data = request.get_json() or {}
    filename = data.get("filename", "").strip()

    filepath = os.path.join(
        phillip_core.DATA_FOLDER,
        filename
    )

    if not os.path.exists(filepath):
        return jsonify({
            "error": "Spreadsheet not found."
        }), 404

    if filename == "Welcome.xlsx":
        return jsonify({
            "error": "Welcome.xlsx cannot be deleted."
        }), 400

    os.remove(filepath)

    if phillip_core.current_file == filepath:
        phillip_core.current_file = os.path.join(
            phillip_core.DATA_FOLDER,
            "Welcome.xlsx"
        )

    return jsonify({
        "status": "success",
        "filename": filename
    })

@app.route("/api/create-spreadsheet", methods=["POST"])
def create_spreadsheet():
    data = request.get_json() or {}
    filename = data.get("filename", "").strip()

    if not filename:
        return jsonify({
            "error": "Please enter a spreadsheet name."
        }), 400

    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    filepath = os.path.join(
        phillip_core.DATA_FOLDER,
        filename
    )

    if os.path.exists(filepath):
        return jsonify({
            "error": "A spreadsheet with that name already exists."
        }), 400

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Account Log"

    template = "Quarter Accounting"

    sheet.append(
        phillip_core.template_columns[template]
    )

    workbook.save(filepath)

    return jsonify({
        "status": "success",
        "filename": filename
    })

@app.route("/api/edit-transaction", methods=["POST"])
def edit_transaction():
    data = request.get_json() or {}

    row_number = data.get("row")
    updates = data.get("updates", {})

    if not row_number:
        return jsonify({
            "error": "Transaction row not provided."
        }), 400

    filepath = phillip_core.current_file

    if not os.path.exists(filepath):
        return jsonify({
            "error": "Spreadsheet not found."
        }), 404

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filepath)
        sheet = workbook["Account Log"]

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        for key, value in updates.items():
            if key in headers:
                column = headers.index(key) + 1
                sheet.cell(
                    row=int(row_number),
                    column=column
                ).value = value

        workbook.save(filepath)

        return jsonify({
            "status": "success",
            "row": row_number
        })

    except Exception as error:
        return jsonify({
            "error": str(error)
        }), 400

@app.route("/api/delete-transaction", methods=["POST"])
def delete_transaction():
    data = request.get_json() or {}
    row_number = data.get("row")

    if not row_number:
        return jsonify({
            "error": "Transaction row not provided."
        }), 400

    filepath = phillip_core.current_file

    if not os.path.exists(filepath):
        return jsonify({
            "error": "Spreadsheet not found."
        }), 404

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filepath)
        sheet = workbook["Account Log"]

        if int(row_number) <= 1 or int(row_number) > sheet.max_row:
            return jsonify({
                "error": "Invalid transaction row."
            }), 400

        sheet.delete_rows(int(row_number), 1)

        workbook.save(filepath)

        return jsonify({
            "status": "success",
            "row": row_number
        })

    except Exception as error:
        return jsonify({
            "error": str(error)
        }), 400

@app.route("/api/spreadsheet/<filename>")
def spreadsheet(filename):
    filepath = os.path.join(
        phillip_core.DATA_FOLDER,
        filename
    )

    if not os.path.exists(filepath):
        return jsonify({
            "error": "Spreadsheet not found."
        }), 404

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filepath, read_only=True)
        sheet = workbook.active

        transactions = max(sheet.max_row - 1, 0)

        return jsonify({
            "filename": filename,
            "last_updated": os.path.getmtime(filepath),
            "transactions": transactions
        })

    except Exception as error:
        return jsonify({
            "error": str(error)
        }), 400

@app.route("/api/search-transactions")
def search_transactions():
    query = request.args.get("query", "").strip().lower()

    if not query:
        return jsonify({
            "transactions": []
        })

    filepath = os.path.join(
        phillip_core.DATA_FOLDER,
        phillip_core.current_file
    )

    if not os.path.exists(filepath):
        return jsonify({
            "error": "Spreadsheet not found."
        }), 404

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filepath, data_only=True)
        sheet = workbook.active

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        transactions = []

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):
            values = [
                "" if value is None else str(value)
                for value in row
            ]

            if query in " ".join(values).lower():
                transactions.append({
                    "row": row_number,
                    "data": dict(
                        zip(headers, values)
                    )
                })

        return jsonify({
            "transactions": transactions
        })

    except Exception as error:
        return jsonify({
            "error": str(error)
        }), 400

if __name__ == "__main__":
    app.run(debug=True)
