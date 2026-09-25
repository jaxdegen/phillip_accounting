import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.chart import PieChart, LineChart, Reference
import os
import subprocess
import ollama
import json
from datetime import datetime

expense_words = [
    "bought",
    "buy",
    "buying",
    "purchased",
    "purchase",
    "paid",
    "pay",
    "paying",
    "spent",
    "spend",
    "spending",
    "invested",
    "investing",
    "investment",
    "obtained",
    "obtaining",
    "obtain",
    "financed",
    "financing",
    "to finance",
    "repurchased",
    "repurchasing",
    "charge",
    "charges",
    "charged",
    "donated",
    "donation",
    "donating"
]

income_words = [
    "got paid",
    "sold",
    "sell",
    "selling",
    "received",
    "receive",
    "receiving",
    "earned",
    "earn",
    "earning",
    "collected",
    "collect",
    "collecting",
    "distribute",
    "distributed",
    "distributing",
    "deposit",
    "depositing",
    "deposited",
    "given",
    "gave me"
]

# The spreadsheet Phillip is currently using
DATA_FOLDER = os.path.expanduser("~/Documents/Phillip Accounting")

os.makedirs(DATA_FOLDER, exist_ok=True)

current_file = os.path.join(DATA_FOLDER, "accounting.xlsx")

ai_name = "Phillip"

template_options = [
    "Quarter Accounting",
   # "Personal Taxes",
   # "Personal Budget",
    "Sales Tracking",
    "Checking",
   # "Custom"
]

template_columns = {

    "Quarter Accounting": [
        "Date",
        "Vendor/Client",
        "Description",
        "Quantity",
        "Category",
        "Type",
        "Amount"
    ],

   # "Personal Taxes": [
   #     "Date",
   #     "Payee",
   #     "Description",
   #     "Tax Category",
   #     "Deductible",
   #     "Amount"
   # ],

   # "Personal Budget": [
   #     "Date",
   #     "Description",
   #     "Category",
   #     "Type",
   #     "Amount"
   # ],

    "Sales Tracking": [
        "Date",
        "Customer",
        "Product",
        "Quantity",
        "Unit Price",
        "Total"
    ],

    "Checking": [
        "Date",
        "Description",
        "Category",
        "Direction",
        "Payment Method",
        "Debit",
        "Credit",
        "Cash",
        "Total"
    ],

   # "Custom": [
   #     "Date",
   #     "Description",
   #     "Category",
   #     "Amount"
   # ]
}

template_ai_instructions = {

    "Quarter Accounting": """
Return exactly these fields:
vendor, description, quantity, category, type, amount

Rules:
- vendor = company, person, or customer involved
- description = what was purchased, sold, or paid for
- quantity = number of items, or 1 if not stated
- category = appropriate accounting category
- type = ONLY "Income" or "Expense"
- amount = numerical amount only
- If a numerical value is not provided, return 0, never "none" or "null"
""",

    "Personal Taxes": """
Return exactly these fields:
payee, description, tax_category, deductible, amount

Rules:
- payee = person or company paid
- description = what the transaction was for
- tax_category = appropriate tax category
- deductible = ONLY true or false
- amount = numerical amount only
- If a numerical value is not provided, return 0, never "none" or "null"
""",

    "Personal Budget": """
Return exactly these fields:
description, category, type, amount

Rules:
- description = what the transaction was for
- category = appropriate personal budget category
- type = ONLY "Income" or "Expense"
- amount = numerical amount only
- If a numerical value is not provided, return 0, never "none" or "null"
""",

    "Sales Tracking": """
Return exactly these fields:
customer, product, quantity, unit_price, total

Rules:
- customer = customer who purchased the product
- product = product or service sold
- quantity = number sold
- unit_price = price per item
- total = total sale amount
- If a numerical value is not provided, return 0, never "none" or "null"
""",

    "Checking": """
Return exactly these fields:
description, category, direction, payment_method, debit, credit, cash

Rules:
- description = description of the transaction
- category = appropriate checking-account category
- direction = ONLY "Income" or "Expense"
- payment method = ONLY "Debit", "Credit" or "Cash"

For Income transactions:
- paycheck, wages, salary, or direct deposit from an employer = "Paycheck"
- money earned from a business = "Business Income"
- interest earned from a bank account = "Interest"
- dividends or investment earnings = "Investment Income"
- tax refunds, returned purchases, or other refunds = "Refund"
- money received as a gift = "Gift"
- transfers from another account = "Transfer"
- Venmo, Cash App, PayPal, Zelle, Apple Pay, Apple Cash, Google Wallet, Google Pay, Samsung Pay, Chime, Revolut, Snapcash, Meta Pay, or Splitwise transactions = "Digital Wallets"
- other unidentified income = "Other Income"

For Expense transactions:
- groceries = "Groceries"
- rent or mortgage = "Housing"
- books, tuition, or school supplies = "Academic Costs"
- utilities = "Utilities"
- restaurants or eating out = "Dining"
- fuel, gas, plane ticket, bus ticket, train ticket, cruise ticket, or ferry ticket = "Transportation"
- hotel fees or car rental = "Vacation"
- purchases of clothing = "Clothing"
- medical or healthcare expenses = "Healthcare"
- purchases online = "Digital Purchases"
- entertainment = "Entertainment"
- subscriptions = "Subscriptions"
- taxes = "Taxes"
- donation = "Donation"
- Venmo, Cash App, PayPal, Zelle, Apple Pay, Apple Cash, Google Wallet, Google Pay, Samsung Pay, Chime, Revolut, Snapcash, Meta Pay, or Splitwise transactions = "Digital Wallets"
- other unidentified expenses = "Other Expense"

- debit = numerical amount involved with a debit card, otherwise 0
- credit = numerical amount involved with a credit card, otherwise 0
- cash = numerical amount involved with cash, otherwise 0
- Do not calculate total
- If a numerical value is not provided, return 0, never "none" or "null"
""",

    "Custom": """
Return exactly these fields:
description, category, amount

Rules:
- description = description of the transaction
- category = appropriate category
- amount = numerical amount only
- If a numerical value is not provided, return 0, never "none" or "null"
"""
}

def create_new_spreadsheet():
    global current_file

    filename = new_file_entry.get().strip()

    if not filename:
        messagebox.showwarning(
            "Phillip",
            "Please enter a name for the spreadsheet."
        )
        return

    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Account Log"

    # Get the selected template
    template = selected_template.get()

    # Get the columns for that template
    columns = template_columns[template]

    # Create the headers
    sheet.append(columns)

    workbook.save(filename)

    # Make the new spreadsheet Phillip's current spreadsheet
    current_file = os.path.abspath(filename)

    messagebox.showinfo(
        "Phillip",
        f"New spreadsheet created:\n\n{filename}"
    )

    new_file_entry.delete(0, tk.END)

    status_label.config(
        text=f"Current spreadsheet: {filename}",
        font=("Times New Roman", 11)
    )

    # Refresh spreadsheet list
    refresh_spreadsheet_list()

def safe_float(value, default=0):
    if value is None:
        return default

    if isinstance(value, str):
        value = value.strip().lower()

        if value in ["none", "null", "n/a", "na", ""]:
            return default

        value = value.replace("$", "").replace(",", "")

    try:
        return float(value)
    except (ValueError, TypeError):
        return default

def detect_payment_method(transaction, ai_method="Cash"):
    text = transaction.lower()

    debit_terms = [
        "debit",
        "debit card",
        "bank card",
        "checking card",
        "greenlight",
        "direct"
    ]

    credit_terms = [
        "credit",
        "credit card",
        "visa",
        "mastercard",
        "amex",
        "american express",
        "discover",
        "check"
    ]

    cash_terms = [
        "cash",
        "in cash",
        "paid cash",
        "cash payment",
        "with change",
        "dollar bills",
        "coins"
    ]

    if any(term in text for term in debit_terms):
        return "Debit"

    if any(term in text for term in credit_terms):
        return "Credit"

    if any(term in text for term in cash_terms):
        return "Cash"

    if ai_method in ["Debit", "Credit", "Cash"]:
        return ai_method

    return "Cash"

def automatic_visualizations(template):
    """
    Phillip automatically maintains the Checking dashboard.

    Checking Dashboard:
    1. Spending by Category pie chart
    2. Debit vs Credit vs Cash pie chart
    3. Running Total line graph with Debit/Credit/Cash activity
    """

    if not current_file or not os.path.exists(current_file):
        print("No current spreadsheet found.")
        return

    try:
        workbook = load_workbook(current_file)

        if "Account Log" not in workbook.sheetnames:
            print("Account Log sheet not found.")
            workbook.close()
            return

        sheet = workbook["Account Log"]

        # Remove the old Dashboard
        if "Dashboard" in workbook.sheetnames:
            del workbook["Dashboard"]

        dashboard = workbook.create_sheet("Dashboard")

        print("===================================")
        print("PHILLIP AUTOMATIC VISUALIZATIONS")
        print("Template:", template)
        print("Spreadsheet:", current_file)
        print("Rows:", sheet.max_row)
        print("===================================")

        # ==================================================
        # CHECKING
        # ==================================================

        if template == "Checking":

            headers = [
                cell.value
                for cell in sheet[1]
            ]

            print("Checking headers:")
            print(headers)

            required_columns = [
                "Date",
                "Description",
                "Category",
                "Direction",
                "Payment Method",
                "Debit",
                "Credit",
                "Cash",
                "Total"
            ]

            if not all(
                column in headers
                for column in required_columns
            ):

                print(
                    "Checking spreadsheet is missing required columns."
                )

                workbook.close()
                return

            # Find columns
            date_index = headers.index("Date") + 1
            category_index = headers.index("Category") + 1
            direction_index = headers.index("Direction") + 1
            payment_method_index = headers.index("Payment Method") + 1
            debit_index = headers.index("Debit") + 1
            credit_index = headers.index("Credit") + 1
            cash_index = headers.index("Cash") + 1
            total_index = headers.index("Total") + 1

            # ==================================================
            # 1. SPENDING BY CATEGORY
            # ==================================================

            category_totals = {}

            for row in range(2, sheet.max_row + 1):

                direction = str(
                    sheet.cell(
                        row=row,
                        column=direction_index
                    ).value or ""
                ).strip()

                category = str(
                    sheet.cell(
                        row=row,
                        column=category_index
                    ).value or "Other Expense"
                ).strip()

                debit = safe_float(
                    sheet.cell(
                        row=row,
                        column=debit_index
                    ).value
                )

                credit = safe_float(
                    sheet.cell(
                        row=row,
                        column=credit_index
                    ).value
                )

                cash = safe_float(
                    sheet.cell(
                        row=row,
                        column=cash_index
                    ).value
                )

                activity = debit + credit + cash

                # Only expenses belong in the spending chart
                if direction == "Expense":

                    if category not in category_totals:
                        category_totals[category] = 0

                    category_totals[category] += activity

            # Put category data on Dashboard
            dashboard["A1"] = "Category"
            dashboard["B1"] = "Amount Spent"

            category_row = 2

            for category, amount in category_totals.items():

                dashboard.cell(
                    row=category_row,
                    column=1
                ).value = category

                dashboard.cell(
                    row=category_row,
                    column=2
                ).value = amount

                dashboard.cell(
                    row=category_row,
                    column=2
                ).number_format = '$#,##0.00'

                category_row += 1

            # Create category pie chart
            if category_totals:

                category_pie = PieChart()

                category_labels = Reference(
                    dashboard,
                    min_col=1,
                    min_row=2,
                    max_row=category_row - 1
                )

                category_data = Reference(
                    dashboard,
                    min_col=2,
                    min_row=1,
                    max_row=category_row - 1
                )

                category_pie.add_data(
                    category_data,
                    titles_from_data=True
                )

                category_pie.set_categories(
                    category_labels
                )

                category_pie.title = "Spending by Category"

                category_pie.height = 8
                category_pie.width = 12

                dashboard.add_chart(
                    category_pie,
                    "D2"
                )

                print("CATEGORY PIE CHART CREATED")

            # ==================================================
            # 2. DEBIT VS CREDIT VS CASH
            # ==================================================

            debit_total = 0
            credit_total = 0
            cash_total = 0

            for row in range(2, sheet.max_row + 1):

                debit_total += safe_float(
                    sheet.cell(
                        row=row,
                        column=debit_index
                    ).value
                )

                credit_total += safe_float(
                    sheet.cell(
                        row=row,
                        column=credit_index
                    ).value
                )

                cash_total += safe_float(
                    sheet.cell(
                        row=row,
                        column=cash_index
                    ).value
                )

            payment_start_row = 1

            dashboard["D1"] = "Payment Method"
            dashboard["E1"] = "Amount"

            dashboard["D2"] = "Debit"
            dashboard["E2"] = debit_total

            dashboard["D3"] = "Credit"
            dashboard["E3"] = credit_total

            dashboard["D4"] = "Cash"
            dashboard["E4"] = cash_total

            for row in range(2, 5):

                dashboard.cell(
                    row=row,
                    column=5
                ).number_format = '$#,##0.00'

            payment_pie = PieChart()

            payment_labels = Reference(
                dashboard,
                min_col=4,
                min_row=2,
                max_row=4
            )

            payment_data = Reference(
                dashboard,
                min_col=5,
                min_row=1,
                max_row=4
            )

            payment_pie.add_data(
                payment_data,
                titles_from_data=True
            )

            payment_pie.set_categories(
                payment_labels
            )

            payment_pie.title = (
                "Debit vs Credit vs Cash"
            )

            payment_pie.height = 8
            payment_pie.width = 12

            dashboard.add_chart(
                payment_pie,
                "L2"
            )

            print("PAYMENT METHOD PIE CHART CREATED")

            # ==================================================
            # 3. TRANSACTION HISTORY
            # ==================================================

            dashboard["A15"] = "Order"
            dashboard["B15"] = "Date"
            dashboard["C15"] = "Debit"
            dashboard["D15"] = "Credit"
            dashboard["E15"] = "Cash"
            dashboard["F15"] = "Total"

            history_row = 16
            transaction_order = 1

            for row in range(2, sheet.max_row + 1):

                dashboard.cell(
                    row=history_row,
                    column=1
                ).value = transaction_order

                dashboard.cell(
                    row=history_row,
                    column=2
                ).value = sheet.cell(
                    row=row,
                    column=date_index
                ).value

                dashboard.cell(
                    row=history_row,
                    column=3
                ).value = safe_float(
                    sheet.cell(
                        row=row,
                        column=debit_index
                    ).value
                )

                dashboard.cell(
                    row=history_row,
                    column=4
                ).value = safe_float(
                    sheet.cell(
                        row=row,
                        column=credit_index
                    ).value
                )

                dashboard.cell(
                    row=history_row,
                    column=5
                ).value = safe_float(
                    sheet.cell(
                        row=row,
                        column=cash_index
                    ).value
                )

                dashboard.cell(
                    row=history_row,
                    column=6
                ).value = safe_float(
                    sheet.cell(
                        row=row,
                        column=total_index
                    ).value
                )

                # Currency formatting
                for column in [3, 4, 5, 6]:

                    dashboard.cell(
                        row=history_row,
                        column=column
                    ).number_format = '$#,##0.00'

                history_row += 1
                transaction_order += 1

            # ==================================================
            # LINE GRAPH
            # ==================================================

            if sheet.max_row >= 2:

                line = LineChart()

                # Debit, Credit, Cash and Total
                line_data = Reference(
                    dashboard,
                    min_col=3,
                    max_col=6,
                    min_row=15,
                    max_row=history_row - 1
                )

                # Transaction order
                transaction_numbers = Reference(
                    dashboard,
                    min_col=1,
                    min_row=16,
                    max_row=history_row - 1
                )

                line.add_data(
                    line_data,
                    titles_from_data=True
                )

                line.set_categories(
                    transaction_numbers
                )

                line.title = (
                    "Running Total and Payment Activity"
                )

                line.y_axis.title = "Amount"
                line.x_axis.title = "Transaction Order"

                line.height = 12
                line.width = 22

                dashboard.add_chart(
                    line,
                    "L15"
                )

                print("TRANSACTION HISTORY LINE GRAPH CREATED")

        # ==================================================
        # DASHBOARD FORMATTING
        # ==================================================

        dashboard.column_dimensions["A"].width = 22
        dashboard.column_dimensions["B"].width = 18
        dashboard.column_dimensions["C"].width = 15
        dashboard.column_dimensions["D"].width = 18
        dashboard.column_dimensions["E"].width = 18
        dashboard.column_dimensions["F"].width = 18
        dashboard.column_dimensions["L"].width = 18

        # ==================================================
        # SAVE
        # ==================================================

        workbook.save(current_file)

        print("Dashboard saved successfully.")

        workbook.close()

        print("===================================")

    except Exception as e:

        print(
            "Automatic visualization error:",
            e
        )
        
def add_transaction():
    global current_file

    transaction = entry.get().strip()

    if not transaction:
        messagebox.showwarning(
            "Phillip",
            "Please enter a transaction."
        )
        return

    template = selected_template.get()

    try:
        print("Phillip is attempting to contact Ollama...")
        print("Ollama model: llama3.2:3b")

        response = ollama.chat(
            model="llama3.2:3b",
            messages=[
                {
                    "role": "system",
                    "content": f"""
You are Phillip, an accounting assistant.

The user has selected the spreadsheet template:
{template}

Analyze the user's transaction according to the rules for that template.

{template_ai_instructions[template]}

Return ONLY valid JSON.

Do not include explanations outside the JSON.
"""
                },
                {
                    "role": "user",
                    "content": transaction
                }
            ],
            format="json"
        )

        ai_text = response["message"]["content"]

        # Convert the JSON response into Python data
        data = json.loads(ai_text)

    except Exception as e:
        messagebox.showerror(
            "Phillip Error",
            f"Phillip could not process the transaction:\n\n{e}"
        )
        return

    # Automatically get today's date
    date = datetime.now().strftime("%m-%d-%Y")

    try:

        # ---------------------------------
        # QUARTER ACCOUNTING
        # ---------------------------------

        if template == "Quarter Accounting":

            vendor = data.get("vendor", "N/A")
            description = data.get("description", "N/A")
            quantity = data.get("quantity", 1)
            category = data.get("category", "Uncategorized")
            amount = safe_float(data.get("amount", 0))

            transaction_lower = transaction.lower()

            if any(word in transaction_lower for word in expense_words):
                transaction_type = "Expense"

            elif any(word in transaction_lower for word in income_words):
                transaction_type = "Income"

            else:
                transaction_type = data.get("type", "Expense")

            if transaction_type not in ["Income", "Expense"]:
                transaction_type = "Expense"

            row = [
                date,
                vendor,
                description,
                quantity,
                category,
                transaction_type,
                amount
            ]

            currency_columns = [7]


        # ---------------------------------
        # PERSONAL TAXES
        # ---------------------------------

        elif template == "Personal Taxes":

            payee = data.get("payee", "N/A")
            description = data.get("description", "N/A")
            tax_category = data.get(
                "tax_category",
                "Uncategorized"
            )
            deductible = data.get("deductible", False)
            amount = safe_float(data.get("amount", 0))

            row = [
                date,
                payee,
                description,
                tax_category,
                deductible,
                amount
            ]

            currency_columns = [6]


        # ---------------------------------
        # PERSONAL BUDGET
        # ---------------------------------

        elif template == "Personal Budget":

            description = data.get("description", "N/A")
            category = data.get(
                "category",
                "Uncategorized"
            )
            transaction_type = data.get(
                "type",
                "Expense"
            )
            amount = safe_float(data.get("amount", 0))

            if transaction_type not in ["Income", "Expense"]:
                transaction_type = "Expense"

            row = [
                date,
                description,
                category,
                transaction_type,
                amount
            ]

            currency_columns = [5]


        # ---------------------------------
        # SALES TRACKING
        # ---------------------------------

        elif template == "Sales Tracking":

            customer = data.get("customer", "N/A")
            product = data.get("product", "N/A")
            quantity = data.get("quantity", 1)

            unit_price = safe_float(data.get("unit_price", 0))
        

            total = safe_float(data.get("total", 0))
            

            row = [
                date,
                customer,
                product,
                quantity,
                unit_price,
                total
            ]

            currency_columns = [5, 6]


        # ---------------------------------
        # CHECKING
        # ---------------------------------

        elif template == "Checking":

            description = data.get(
                "description",
                "N/A"
            )

            category = data.get(
                "category",
                "Uncategorized"
            )

            direction = data.get(
                "direction",
                "expense"
            )

            payment_method = detect_payment_method(
                transaction,
                data.get("payment_method", "Cash")
            )

            debit = safe_float(
                data.get("debit", 0)
            )

            credit = safe_float(
                data.get("credit", 0)
            )

            cash = safe_float(
                data.get("cash", 0)
            )

            activity = debit + credit + cash

            if payment_method == "Debit":
                debit = activity
                credit = 0
                cash = 0

            elif payment_method == "Credit":
                credit = activity
                debit = 0
                cash = 0

            elif payment_method == "Cash":
                cash = activity
                debit = 0
                credit = 0
                        
            if direction not in ["Income", "Expense"]:
                direction = "Expense"

            # ---------------------------------
            # CALCULATE TRANSACTION ACTIVITY
            # ---------------------------------

            activity = debit + credit + cash

            # ---------------------------------
            # GET PREVIOUS TOTAL
            # ---------------------------------

            if os.path.exists(current_file):

                existing_workbook = load_workbook(current_file)
                existing_sheet = existing_workbook["Account Log"]

                if existing_sheet.max_row > 1:

                    previous_total = safe_float(
                        existing_sheet.cell(
                            row=existing_sheet.max_row,
                            column=9
                        ).value
                    )

                else:

                    previous_total = 0

                existing_workbook.close()

            else:

                previous_total = 0

            # ---------------------------------
            # UPDATE RUNNING TOTAL
            # ---------------------------------

            if direction == "Income":

                total = previous_total + activity

            else:

                total = previous_total - activity

            # ---------------------------------
            # CREATE CHECKING ROW
            # ---------------------------------


            row = [
                date,
                description,
                category,
                direction,
                payment_method,
                debit,
                credit,
                cash,
                total
            ]

            currency_columns = [6, 7, 8, 9]


        # ---------------------------------
        # CUSTOM
        # ---------------------------------

        elif template == "Custom":

            description = data.get(
                "description",
                "N/A"
            )

            category = data.get(
                "category",
                "Uncategorized"
            )

            amount = safe_float(data.get("amount", 0))
            

            row = [
                date,
                description,
                category,
                amount
            ]

            currency_columns = [4]


        else:

            raise ValueError(
                f"Unknown spreadsheet template: {template}"
            )

    except Exception as e:

        messagebox.showerror(
            "Phillip Error",
            f"Phillip could not organize the transaction:\n\n{e}"
        )

        return


    # ---------------------------------
    # OPEN THE SPREADSHEET
    # ---------------------------------

    if os.path.exists(current_file):

        workbook = load_workbook(current_file)
        sheet = workbook["Account Log"]

    else:

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Account Log"

        # Create the correct headers
        sheet.append(
            template_columns[template]
        )


    # ---------------------------------
    # ADD THE TRANSACTION
    # ---------------------------------

    sheet.append(row)


    # ---------------------------------
    # FORMAT CURRENCY CELLS
    # ---------------------------------

    for column in currency_columns:

        amount_cell = sheet.cell(
            row=sheet.max_row,
            column=column
        )

        amount_cell.number_format = '$#,##0.00'


    # Save the spreadsheet
    workbook.save(current_file)

    print("ABOUT TO CREATE VISUALIZATIONS")
          
    # Phillip automatically analyzes the updated
    # spreadsheet and maintains its visualizations.
    automatic_visualizations(template)

    print("FINISHED CREATING VISUALIZATIONS")

    # Open the spreadsheet
    subprocess.run([
        "open",
        "-a",
        "Numbers",
        current_file
    ])


    status_label.config(
        text=f"Data successfully added to {os.path.basename(current_file)}!",
        font=("Times New Roman", 11)
    )

    entry.delete(0, tk.END)
    
# -------------------------
# SEARCH SPREADSHEETS
# -------------------------

def refresh_spreadsheet_list():
    spreadsheet_list.delete(0, tk.END)

    search_text = search_entry.get().lower()

    folder = os.path.dirname(os.path.abspath(__file__))

    for filename in os.listdir(folder):
        if filename.endswith(".xlsx"):
            if search_text in filename.lower():
                spreadsheet_list.insert(tk.END, filename)


def search_spreadsheets(event=None):
    refresh_spreadsheet_list()


def select_spreadsheet():
    global current_file

    selection = spreadsheet_list.curselection()

    if not selection:
        messagebox.showwarning(
            "Phillip",
            "Please select a spreadsheet."
        )
        return

    filename = spreadsheet_list.get(selection[0])

    current_file = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        filename
    )

    status_label.config(
        text=f"Current spreadsheet: {filename}",
        font=("Times New Roman", 11)
    )

def show_help():
    help_window = tk.Toplevel(window)
    help_window.title("Phillip Accounting - Help")
    help_window.geometry("600x600")

    title = tk.Label(
        help_window,
        text="PHILLIP ACCOUNTING - HELP",
        font=("Times New Roman", 22, "bold")
    )
    title.pack(pady=20)

    instructions = """
Getting Started:

1. Select a spreadsheet template.
2. Enter your accounting information naturally.
3. Click "Enter Data".
4. Phillip will interpret the transaction and add it
   to the current spreadsheet.



Spreadsheet Templates:

(Select Template Before Creating New Spreadsheet or Inputing Data)

    Quarter Accounting:
        General accounting records.

    Sales Tracking:
        Tracks customers, products, quantities and sales.

    Checking:
        Tracks income, expenses, payment methods and running balance.
        


Managing Spreadsheets:

    Use "Search Existing Spreadsheets" to find spreadsheets.

    Select a spreadsheet and click "Select Spreadsheet"
    to make it the current spreadsheet.

    Use "Create New Spreadsheet" to create a new workbook.

    Use "Delete Spreadsheet" to remove a spreadsheet.



Transaction Examples:

    "I spent $45 at Walmart with my debit card."

    "I received $2,000 from my job."

    "I paid $30 cash for groceries."

-----------------------------------------------------------------------

Phillip will use the information you provide to
categorize and record the transaction.

ALWAYS review automatically generated accounting
information for accuracy.
"""

    text = tk.Text(
        help_window,
        font=("Times New Roman", 13),
        wrap="word"
    )

    text.insert("1.0", instructions)
    text.config(state="disabled")
    text.pack(fill="both", expand=True, padx=20, pady=10)

def delete_spreadsheet():
    global current_file

    selection = spreadsheet_list.curselection()

    if not selection:
        messagebox.showwarning(
            "Phillip",
            "Please select a spreadsheet to delete."
        )
        return

    filename = spreadsheet_list.get(selection[0])

    confirm = messagebox.askyesno(
        "Delete Spreadsheet",
        f"Are you sure you want to permanently delete:\n\n{filename}?"
    )

    if not confirm:
        return

    filepath = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        filename
    )

    try:
        os.remove(filepath)

        # If the deleted spreadsheet was the current spreadsheet,
        # switch back to the default spreadsheet name
        if os.path.abspath(current_file) == os.path.abspath(filepath):
            current_file = "accounting.xlsx"

        refresh_spreadsheet_list()

        status_label.config(
            text=f"Deleted spreadsheet: {filename}",
            font=("Times New Roman", 11)
        )

    except Exception as e:
        messagebox.showerror(
            "Phillip Error",
            f"Phillip could not delete the spreadsheet:\n\n{e}"
        )

# -------------------------
# WINDOW
# -------------------------

window = tk.Tk()

window.title("Phillip Accounting")
window.geometry("640x660")

help_button = tk.Canvas(
    window,
    width=30,
    height=30,
    bg="#1e1e1e",
    highlightthickness=0
)

help_button.place(relx=0.99, rely=0.01, anchor="ne")

help_button.create_oval(
    3,
    3,
    30,
    30,
    fill="#0000FF",
    outline="#000080",
    width=2
)

help_button.create_text(
    16,
    16,
    text="?",
    font=("Times New Roman", 17),
    fill="white"
)

help_button.bind(
    "<Button-1>",
    lambda event: show_help()
)

title = tk.Label(
    window,
    text="PHILLIP ACCOUNTING",
    font=("Times New Roman", 23, "bold")
)

title.pack(pady=18)


selected_template = tk.StringVar(master=window)
selected_template.set("Quarter Accounting")

template_label = tk.Label(
    window,
    text="Select Spreadsheet Template:",
    font=("Times New Roman", 12)
)

template_label.pack(pady=(15, 5))


template_menu = tk.OptionMenu(
    window,
    selected_template,
    *template_options
)

template_menu.config(
    font=("Times New Roman", 11),
    width=25
)

template_menu.pack(pady=5)


instructions = tk.Label(
    window,
    text="Input New Data Here:",
    font=("Times New Roman", 15)
)

instructions.pack()


entry = tk.Entry(
    window,
    width=45,
    font=("Courier New", 12),
    bg="#F8F6F0",
    fg="black",
    highlightbackground="#808080",
    highlightcolor="#808080",    
)

entry.pack(pady=5)


submit_button = tk.Button(
    window,
    text="Enter Data",
    font=("Times New Roman", 13, "bold"),
    bg="#FFD700",
    fg="black",
    activebackground="#FFD700",
    activeforeground="black",
    highlightbackground="#FFD700",
    highlightcolor="#FFD700",    
    command=add_transaction
)

submit_button.pack(pady=3)


# -------------------------
# SPREADSHEET SEARCH
# -------------------------

search_label = tk.Label(
    window,
    text="Search Existing Spreadsheets:",
    font=("Times New Roman", 11)
)

search_label.pack(pady=(15, 5))


search_entry = tk.Entry(
    window,
    width=40,
    font=("Courier New", 12),
    bg="#F8F6F0",
    fg="black",
    highlightbackground="#808080",
    highlightcolor="#808080",    
)

search_entry.pack(pady=3)

search_entry.bind(
    "<KeyRelease>",
    search_spreadsheets
)


spreadsheet_list = tk.Listbox(
    window,
    width=35,
    height=3,
    font=("Courier New", 11)
)

spreadsheet_list.pack(pady=2)


select_button = tk.Button(
    window,
    text="Select Spreadsheet",
    font=("Times New Roman", 11),
    bg="white",
    fg="black",
    command=select_spreadsheet
)

select_button.pack(pady=2)

delete_button = tk.Button(
    window,
    text="Delete Spreadsheet",
    font=("Times New Roman", 11),
    bg="#B71C1C",
    fg="black",
    activebackground="#8B0000",
    activeforeground="black",
    highlightbackground="#B71C1C",
    highlightcolor="#B71C1C",
    command=delete_spreadsheet
)

delete_button.pack(pady=0)

# -------------------------
# NEW SPREADSHEET
# -------------------------

status_label = tk.Label(
    window,
    text=f"Current spreadsheet: {current_file}",
    font=("Times New Roman", 11)
)

status_label.pack(pady=10)


new_file_entry = tk.Entry(
    window,
    width=40,
    font=("Courier New", 12),
    bg="#F8F6F0",
    fg="black",
    highlightbackground="#808080",
    highlightcolor="#808080",    
)

new_file_entry.pack(pady=5)


new_file_button = tk.Button(
    window,
    text="Create New Spreadsheet",
    font=("Courier New", 12),
    command=create_new_spreadsheet
)

new_file_button.pack(pady=5)

# -------------------------
# FOOTER
# -------------------------

footer = tk.Label(
    window,
    text="Powered by DeBeaul Tech and Solutions",
    font=("Times New Roman", 10)
)

footer.pack(pady=5)

footer = tk.Label(
    window,
    text="Version 8.05",
    font=("Arial", 8)
)

footer.pack(pady=0)

footer = tk.Label(
    window,
    text="Base Edition",
    font=("Arial", 9)
)

footer.pack(pady=0)

# Load spreadsheets when Phillip starts
refresh_spreadsheet_list()


window.mainloop()
